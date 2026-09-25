import os
from datetime import timedelta

from openpyxl import load_workbook

from app.controllers.reports import generate_pa_report_excel, generate_pa_report_pdf
from app.models import Event, User, db
from app.workflow import utcnow


def test_report_generation_is_isolated_and_preserves_scoring_rows(app, tmp_path, user_factory):
    reports_folder = tmp_path / 'reports'
    app.config['REPORTS_FOLDER'] = str(reports_folder)
    with app.app_context():
        organizer = user_factory('organizer', 'organizer@example.com')
        admin = user_factory('admin', 'admin@example.com')
        db.session.add_all([organizer, admin])
        db.session.commit()
        event = Event(
            title='=Research & Expo',
            description='Demo event',
            venue='Auditorium',
            event_date=utcnow() - timedelta(days=1),
            registration_deadline=utcnow() - timedelta(days=2),
            max_participants=40,
            is_intercollege=False,
            organizer_id=organizer.id,
            status='completed',
        )
        db.session.add(event)
        db.session.commit()

        pdf_path = generate_pa_report_pdf(event, academic_year='2026-2027', generated_by_id=admin.id)
        xlsx_path = generate_pa_report_excel(event, academic_year='2026-2027', generated_by_id=admin.id)

        assert os.path.exists(pdf_path)
        assert os.path.exists(xlsx_path)
        assert pdf_path.endswith('.pdf')
        assert xlsx_path.endswith('.xlsx')
        assert str(reports_folder) in pdf_path

        workbook = load_workbook(xlsx_path, data_only=False)
        worksheet = workbook['PA Report']
        labels = [worksheet.cell(row=row, column=1).value for row in range(1, worksheet.max_row + 1)]
        assert 'Bonus Feedback' in labels
        assert 'Total PA Summary' in labels
        assert worksheet['B4'].data_type == 's'


def test_report_route_requires_completed_event(client, app, user_factory):
    with app.app_context():
        admin = user_factory('admin', 'admin@example.com')
        organizer = user_factory('organizer', 'organizer@example.com')
        db.session.add_all([admin, organizer])
        db.session.commit()
        event = Event(
            title='Pending report event',
            venue='Auditorium',
            event_date=utcnow() + timedelta(days=2),
            registration_deadline=utcnow() + timedelta(days=1),
            max_participants=40,
            organizer_id=organizer.id,
            status='pending_hod',
        )
        db.session.add(event)
        db.session.commit()
        event_id = event.id

    client.post('/login', data={'email': 'admin@example.com', 'password': 'password123'})
    response = client.post(
        f'/reports/export/{event_id}',
        data={'format': 'pdf', 'academic_year': '2026-2027'},
    )
    assert response.status_code == 302
    assert response.headers['Location'].endswith('/admin/dashboard')
