import os
import re
import time
import uuid
from pathlib import Path
from xml.sax.saxutils import escape

from flask import Blueprint, current_app, flash, redirect, request, send_file, url_for
from flask_login import current_user
from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.controllers.auth import role_required
from app.controllers.admin import calculate_pa_points
from app.models import Event, ReportRecord, db
from app.workflow import current_academic_year, utcnow


reports_bp = Blueprint('reports', __name__)
ACADEMIC_YEAR_PATTERN = re.compile(r'^\d{4}-\d{4}$')


def _ensure_upload_dir():
    folder = current_app.config.get('REPORTS_FOLDER', 'static/uploads')
    folder = os.path.abspath(folder)
    os.makedirs(folder, exist_ok=True)
    return folder


def _valid_academic_year(value):
    if not ACADEMIC_YEAR_PATTERN.fullmatch(str(value or '')):
        return False
    start, end = (int(part) for part in value.split('-'))
    return end == start + 1


def _cleanup_old_reports(folder):
    """Remove only generated report artifacts according to configured limits."""
    now = time.time()
    retention_days = max(0, int(current_app.config.get('REPORT_RETENTION_DAYS', 30)))
    max_files = max(1, int(current_app.config.get('MAX_REPORT_FILES', 100)))
    paths = [path for path in Path(folder).glob('pa_report_*') if path.is_file()]

    if retention_days:
        cutoff = now - retention_days * 86400
        for path in paths:
            try:
                if path.stat().st_mtime < cutoff:
                    path.unlink()
            except OSError:
                current_app.logger.warning('Could not remove old report %s', path)

    remaining = sorted(
        (path for path in Path(folder).glob('pa_report_*') if path.is_file()),
        key=lambda path: path.stat().st_mtime,
        reverse=True,
    )
    for path in remaining[max_files:]:
        try:
            path.unlink()
        except OSError:
            current_app.logger.warning('Could not enforce report retention for %s', path)


def _new_report_path(folder, event_id, extension):
    timestamp = utcnow().strftime('%Y%m%d%H%M%S%f')
    file_name = f'pa_report_{event_id}_{timestamp}_{uuid.uuid4().hex[:8]}.{extension}'
    return os.path.join(folder, file_name)


def _record_report(file_path, academic_year, generated_by_id):
    if not generated_by_id:
        return
    db.session.add(
        ReportRecord(
            generated_by=generated_by_id,
            academic_year=academic_year,
            file_path=file_path,
        )
    )
    db.session.commit()


def _paragraph(text, style):
    return Paragraph(escape(str(text)), style)


def generate_pa_report_pdf(event, academic_year=None, generated_by_id=None):
    academic_year = academic_year or current_academic_year()
    if not _valid_academic_year(academic_year):
        raise ValueError('Academic year must use the format YYYY-YYYY.')

    folder = _ensure_upload_dir()
    _cleanup_old_reports(folder)
    metrics = calculate_pa_points(event, commit=False)
    file_path = _new_report_path(folder, event.id, 'pdf')

    doc = SimpleDocTemplate(
        file_path,
        pagesize=letter,
        leftMargin=40,
        rightMargin=40,
        topMargin=40,
        bottomMargin=40,
    )
    styles = getSampleStyleSheet()
    story = [
        _paragraph('CampusPulse', styles['Title']),
        _paragraph(f'Academic Year: {academic_year}', styles['Heading2']),
        Spacer(1, 12),
        _paragraph(f'Event: {event.title}', styles['Heading2']),
        _paragraph(f'Organizer: {event.organizer.name}', styles['BodyText']),
        _paragraph(f'Venue: {event.venue}', styles['BodyText']),
        _paragraph(f'Date: {event.event_date.date()}', styles['BodyText']),
        Spacer(1, 12),
    ]

    table_data = [
        ['Metric', 'Value'],
        ['Base Points', metrics['base']],
        ['Bonus Participants', metrics['bonus_participants']],
        ['Bonus Feedback', metrics['bonus_feedback']],
        ['Bonus Intercollege', metrics['bonus_intercollege']],
        ['Total PA Points', metrics['total']],
        ['Participants', metrics['participant_count']],
        ['Average Rating', f"{metrics['average_rating']:.2f}"],
    ]
    table = Table(table_data, colWidths=[220, 200])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1f3a5f')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('GRID', (0, 0), (-1, -1), 1, colors.black),
        ('ALIGN', (1, 1), (-1, -1), 'RIGHT'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.whitesmoke, colors.white]),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
    ]))
    story.extend([table, Spacer(1, 18), _paragraph(f'Total PA Summary: {metrics["total"]} points', styles['Heading2'])])
    doc.build(story)
    _record_report(file_path, academic_year, generated_by_id)
    return file_path


def _set_excel_text(cell, value):
    cell.value = str(value)
    if cell.value[:1] in {'=', '+', '-', '@'}:
        cell.data_type = 's'


def _set_excel_number(cell, value):
    cell.value = value
    cell.data_type = 'n'


def generate_pa_report_excel(event, academic_year=None, generated_by_id=None):
    academic_year = academic_year or current_academic_year()
    if not _valid_academic_year(academic_year):
        raise ValueError('Academic year must use the format YYYY-YYYY.')

    folder = _ensure_upload_dir()
    _cleanup_old_reports(folder)
    metrics = calculate_pa_points(event, commit=False)
    file_path = _new_report_path(folder, event.id, 'xlsx')

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = 'PA Report'
    _set_excel_text(worksheet['A1'], 'CampusPulse')
    worksheet['A1'].font = Font(size=14, bold=True)

    metadata = (
        ('Academic Year', academic_year),
        ('Event', event.title),
        ('Venue', event.venue),
        ('Organizer', event.organizer.name),
    )
    for row_index, (label, value) in enumerate(metadata, start=3):
        _set_excel_text(worksheet.cell(row=row_index, column=1), label)
        _set_excel_text(worksheet.cell(row=row_index, column=2), value)

    rows = [
        ('Metric', 'Value'),
        ('Base Points', metrics['base']),
        ('Bonus Participants', metrics['bonus_participants']),
        ('Bonus Feedback', metrics['bonus_feedback']),
        ('Bonus Intercollege', metrics['bonus_intercollege']),
        ('Total PA Points', metrics['total']),
        ('Participants', metrics['participant_count']),
        ('Average Rating', metrics['average_rating']),
    ]
    for row_index, (label, value) in enumerate(rows, start=8):
        _set_excel_text(worksheet.cell(row=row_index, column=1), label)
        if isinstance(value, (int, float)):
            _set_excel_number(worksheet.cell(row=row_index, column=2), value)
        else:
            _set_excel_text(worksheet.cell(row=row_index, column=2), value)

    summary_row = worksheet.max_row + 2
    _set_excel_text(worksheet.cell(row=summary_row, column=1), 'Total PA Summary')
    _set_excel_number(worksheet.cell(row=summary_row, column=2), metrics['total'])
    worksheet.cell(row=summary_row, column=1).font = Font(bold=True)
    worksheet.cell(row=summary_row, column=2).font = Font(bold=True)
    worksheet.column_dimensions['A'].width = 25
    worksheet.column_dimensions['B'].width = 28
    workbook.save(file_path)
    _record_report(file_path, academic_year, generated_by_id)
    return file_path


@reports_bp.route('/export/<int:event_id>', methods=['POST'])
@role_required('admin')
def export_report(event_id):
    event = db.get_or_404(Event, event_id)
    if event.status != 'completed':
        flash('Reports can only be generated for completed events.', 'danger')
        return redirect(url_for('admin.dashboard'))

    export_format = request.form.get('format', 'pdf').strip().lower()
    academic_year = request.form.get('academic_year', '').strip() or current_academic_year()
    if export_format not in {'pdf', 'excel'}:
        flash('Choose a valid report format.', 'danger')
        return redirect(url_for('admin.dashboard'))
    if not _valid_academic_year(academic_year):
        flash('Academic year must use the format YYYY-YYYY.', 'danger')
        return redirect(url_for('admin.dashboard'))

    try:
        if export_format == 'excel':
            file_path = generate_pa_report_excel(event, academic_year, generated_by_id=current_user.id)
        else:
            file_path = generate_pa_report_pdf(event, academic_year, generated_by_id=current_user.id)
    except (OSError, ValueError):
        db.session.rollback()
        flash('The report could not be generated. Please try again.', 'danger')
        return redirect(url_for('admin.dashboard'))

    return send_file(
        file_path,
        as_attachment=True,
        download_name=os.path.basename(file_path),
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet' if export_format == 'excel' else 'application/pdf',
    )
